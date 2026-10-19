"""
07 — Monta a planilha ENXUTA do primeiro modelo: "4. Dados Gerdau/Database Valuation Gerdau.xlsx".

É esta planilha que vai para o GitHub e que o Colab lê. Ela tem só o necessário para o FCD
em 5 datas-base anuais (uma por exercício de 2020 a 2024, logo após a divulgação da DFP):

  Datas-base          as 5 datas, o preço da GGBR4 e o valor de mercado em cada uma
  Fundamentos         DRE, fluxo de caixa e balanço anuais da Gerdau consolidada, 2009-2024
  Preços diários      pregões desde 2009: GGBR4, Ibovespa e CDI (para o beta e a estratégia)
  Macro anual         PIB, IPCA, CDI, Selic e câmbio
  Expectativas Focus  mediana do Focus disponível em cada data-base (PIB, IPCA, Selic)
  Mercado nas datas   spread de crédito Baa, EMBI+ Brasil e % da receita no Brasil (custo da dívida e risco-país)
  Retornos EUA (anual) retornos totais reais de ações e títulos de 10 anos dos EUA (prêmio maduro, Shiller)

Os cabeçalhos já trazem nome legível e unidade; o Colab lê com pd.read_excel(url, sheet_name="Fundamentos").

Lê a base completa gerada pelo 06 ("Tabelas intermediárias (CSV)/Base completa (arquivo).xlsx") e os CSVs do 04b;
não baixa nada.
Valores monetários em R$ milhões; quantidade de ações em milhões.
"""

import os
import site
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))  # pasta dos scripts (código próprio)
sys.path.append(site.getusersitepackages())  # pacotes instalados com pip --user (o -I os desativa)
import pandas as pd  # noqa: E402
from _comum import PLANILHA, PLANILHA_COMPLETA  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

EXERCICIOS_BASE = [2020, 2021, 2022, 2023, 2024]  # uma data-base por exercício
INICIO_HISTORICO = 2009  # janelas longas de estimação começam aqui

base = pd.read_excel(PLANILHA_COMPLETA, sheet_name=None)

# ------------------------------------------------------------------ Fundamentos (R$ mil -> R$ milhões)
# Só exercícios até o último usado em data-base: nada de informação posterior entra na planilha.
CONTAS = [
    "receita_liquida",
    "cpv",
    "lucro_bruto",
    "ebit",
    "impairment",
    "ebit_ex_impairment",
    "da",
    "ebitda",
    "despesas_financeiras",
    "lair",
    "ir_cs",
    "lucro_liquido",
    "lucro_controladores",
    "capex",
    "contas_receber",
    "estoques",
    "fornecedores",
    "caixa_aplicacoes",
    "divida_bruta",
    "arrendamentos_total",
    "divida_liquida",
    "pl_total",
    "pl_nao_controladores",
]
f = base["Fundamentos_anuais"]
f = f[(f["exercicio"] >= INICIO_HISTORICO) & (f["exercicio"] <= max(EXERCICIOS_BASE))].copy()
fund = f[["exercicio", "dt_receb_primeira"]].rename(columns={"dt_receb_primeira": "data_divulgacao"})
for c in CONTAS:
    fund[c] = f[c] / 1e3
fund["acoes_circulacao"] = f["acoes_circulacao"] / 1e6
fund["data_divulgacao"] = pd.to_datetime(fund["data_divulgacao"]).dt.date

# ------------------------------------------------------------------ Preços diários (só pregões da B3)
p = base["Precos_diarios"]
p = p[(p["pregao_b3"] == 1) & (p["data"] >= f"{INICIO_HISTORICO}-01-01")]
m = base["Macro_diario"][["data", "cdi_ad"]]
precos = (
    p[["data", "ggbr4_fech_nominal", "ggbr4_ajust", "ibov"]]
    .merge(m, on="data", how="left")
    .rename(columns={"ggbr4_fech_nominal": "ggbr4_nominal", "ggbr4_ajust": "ggbr4_ajustado", "cdi_ad": "cdi_diario"})
)
precos = precos.dropna(subset=["ibov"])  # 1 pregão sem Ibovespa no Yahoo
precos["data"] = pd.to_datetime(precos["data"]).dt.date

# ------------------------------------------------------------------ Datas-base
# Data-base = 1º pregão estritamente depois da 1ª entrega da DFP na CVM.
linhas = []
for ex in EXERCICIOS_BASE:
    r = fund[fund["exercicio"] == ex].iloc[0]
    dia = precos[precos["data"] > r["data_divulgacao"]].iloc[0]
    linhas.append(
        {
            "data_base": dia["data"],
            "exercicio_disponivel": ex,
            "data_divulgacao_dfp": r["data_divulgacao"],
            "ggbr4_preco": dia["ggbr4_nominal"],
            "acoes_circulacao": r["acoes_circulacao"],
            "valor_mercado": dia["ggbr4_nominal"] * r["acoes_circulacao"],
            "divida_liquida": r["divida_liquida"],
        }
    )
datas = pd.DataFrame(linhas)

# ------------------------------------------------------------------ Macro anual
ma = base["Macro_anual"]
macro = ma[(ma["ano"] >= INICIO_HISTORICO) & (ma["ano"] <= max(EXERCICIOS_BASE))][
    ["ano", "pib_real_var", "ipca_ano", "cdi_ano", "selic_meta_fim", "ptax_fim"]
]

# ------------------------------------------------------------------ Expectativas (Focus) em cada data-base
# Última pesquisa mensal publicada ANTES da data-base. t0 = ano da pesquisa, t1..t4 = anos seguintes.
mm = base["Macro_mensal"].copy()
mm["focus_data_pesquisa"] = pd.to_datetime(mm["focus_data_pesquisa"]).dt.date
cols_focus = [c for c in mm.columns if c.startswith(("focus_pib_t", "focus_ipca_t", "focus_selic_t"))]
exp = []
for d in datas["data_base"]:
    ult = mm[mm["focus_data_pesquisa"] < d].sort_values("focus_data_pesquisa").iloc[-1]
    linha = {"data_base": d, "data_pesquisa": ult["focus_data_pesquisa"], "ano_t0": ult["focus_data_pesquisa"].year}
    linha.update({c.replace("focus_", ""): ult[c] for c in cols_focus})
    exp.append(linha)
expect = pd.DataFrame(exp).dropna(axis=1, how="all")  # o Focus ainda não publicava t4 nessas datas

# ------------------------------------------------------------------ Mercado nas datas (custo de capital)
# Gerado pelo 04b. Só o que era público antes da data-base:
#   spread Baa − Treasury de 10 anos: média do mês anterior à data-base (FRED BAA e GS10);
#   EMBI+ Brasil: último dia útil antes da data-base (Ipeadata). A série acabou em 30/07/2024: a data de
#     fev/2025 usa esse último valor (o mais recente público na data);
#   receita no Brasil: nota geográfica do 20-F do exercício, como publicada pela 1ª vez
#     (ver "Brasil x exterior e câmbio (20-F)/Relatório Brasil x exterior.md").
spread = pd.read_csv(
    os.path.join(os.path.dirname(PLANILHA_COMPLETA), "spread_baa_mensal.csv"), parse_dates=["mes"]
).set_index("mes")["spread_baa"]
embi = pd.read_csv(os.path.join(os.path.dirname(PLANILHA_COMPLETA), "embi_diario.csv"), parse_dates=["data"]).set_index(
    "data"
)["embi_pb"]
RECEITA_BRASIL = {2020: 42.9, 2021: 48.2, 2022: 44.2, 2023: 40.2, 2024: 43.7}  # % = 100 − % do exterior
merc = []
for _, r in datas.iterrows():
    d = pd.Timestamp(r["data_base"])
    mes_anterior = (d - pd.DateOffset(months=1)).to_period("M").to_timestamp()
    e = embi[embi.index < d]
    merc.append(
        {
            "data_base": r["data_base"],
            "spread_baa": spread[spread.index <= mes_anterior].iloc[-1],
            "embi": e.iloc[-1] / 100,
            "data_embi": e.index[-1].date(),
            "receita_brasil": RECEITA_BRASIL[r["exercicio_disponivel"]],
        }
    )
mercado = pd.DataFrame(merc)

# Retornos totais reais dos EUA em dezembro de cada ano (Shiller), para o prêmio maduro histórico.
# Só até o último exercício usado: o notebook corta em cada data no ano anterior a ela.
eua = pd.read_csv(os.path.join(os.path.dirname(PLANILHA_COMPLETA), "retornos_eua_dezembro.csv"))
eua = eua[eua["ano"] <= max(EXERCICIOS_BASE)]

# ------------------------------------------------------------------ cabeçalhos legíveis (nome + unidade)
mercado = mercado.rename(
    columns={
        "data_base": "Data-base",
        "spread_baa": "Spread Baa − Treasury 10 anos, mês anterior (%)",
        "embi": "EMBI+ Brasil, véspera (%)",
        "data_embi": "Data do EMBI usado",
        "receita_brasil": "Receita no Brasil, último exercício (%)",
    }
)
eua = eua.rename(
    columns={
        "ano": "Ano",
        "acoes_tr_real": "Ações EUA: índice de retorno total real em dez",
        "titulos_tr_real": "Títulos 10 anos EUA: índice de retorno total real em dez",
    }
)
datas = datas.rename(
    columns={
        "data_base": "Data-base",
        "exercicio_disponivel": "Último exercício publicado",
        "data_divulgacao_dfp": "Divulgação da DFP",
        "ggbr4_preco": "GGBR4 na data (R$)",
        "acoes_circulacao": "Ações em circulação (mi)",
        "valor_mercado": "Valor de mercado (R$ mi)",
        "divida_liquida": "Dívida líquida (R$ mi)",
    }
)
fund = fund.rename(
    columns={
        "exercicio": "Ano",
        "data_divulgacao": "Divulgação da DFP",
        "receita_liquida": "Receita líquida",
        "cpv": "CPV",
        "lucro_bruto": "Lucro bruto",
        "ebit": "EBIT (com impairment)",
        "impairment": "Impairment",
        "ebit_ex_impairment": "EBIT sem impairment",
        "da": "Depreciação e amortização",
        "ebitda": "EBITDA",
        "despesas_financeiras": "Despesas financeiras",
        "lair": "Lucro antes do IR",
        "ir_cs": "IR e CSLL",
        "lucro_liquido": "Lucro líquido",
        "lucro_controladores": "Lucro dos controladores",
        "capex": "Capex",
        "contas_receber": "Contas a receber",
        "estoques": "Estoques",
        "fornecedores": "Fornecedores",
        "caixa_aplicacoes": "Caixa e aplicações",
        "divida_bruta": "Dívida bruta (sem arrendamentos)",
        "arrendamentos_total": "Arrendamentos (IFRS 16)",
        "divida_liquida": "Dívida líquida (sem arrendamentos)",
        "pl_total": "Patrimônio líquido",
        "pl_nao_controladores": "PL de não controladores",
        "acoes_circulacao": "Ações em circulação (mi)",
    }
)
precos = precos.rename(
    columns={
        "data": "Data",
        "ggbr4_nominal": "GGBR4 nominal (R$)",
        "ggbr4_ajustado": "GGBR4 ajustada (R$)",
        "ibov": "Ibovespa (pontos)",
        "cdi_diario": "CDI (% ao dia)",
    }
)
macro = macro.rename(
    columns={
        "ano": "Ano",
        "pib_real_var": "PIB real (% a.a.)",
        "ipca_ano": "IPCA (% a.a.)",
        "cdi_ano": "CDI (% a.a.)",
        "selic_meta_fim": "Selic meta em dez (% a.a.)",
        "ptax_fim": "Dólar PTAX em dez (R$/US$)",
    }
)
expect = expect.rename(
    columns={"data_base": "Data-base", "data_pesquisa": "Pesquisa Focus usada", "ano_t0": "Ano t (ano da pesquisa)"}
)
nomes_focus = {"pib": "PIB", "ipca": "IPCA", "selic": "Selic"}
expect = expect.rename(
    columns=lambda c: c if "_t" not in c else f"{nomes_focus[c.split('_t')[0]]} ano t+{c.split('_t')[1]} (%)"
)

# ------------------------------------------------------------------ grava com formatação simples
abas = {
    "Datas-base": datas,
    "Fundamentos": fund,
    "Preços diários": precos,
    "Macro anual": macro,
    "Expectativas Focus": expect,
    "Mercado nas datas": mercado,
    "Retornos EUA (anual)": eua,
}
CAB_FUNDO = PatternFill("solid", fgColor="DCE6F1")
CAB_FONTE = Font(bold=True, color="1F3864")
CAB_ALIN = Alignment(wrap_text=True, horizontal="center", vertical="center")
CAB_BORDA = Border(bottom=Side(style="thin", color="8EA9DB"))
CENTRO = Alignment(horizontal="center")


def formato(col, serie):
    """Formato de exibição pela natureza da coluna."""
    if len(serie) and hasattr(serie.iloc[0], "year"):
        return "dd/mm/yyyy"
    if col.startswith("Ano") or col == "Último exercício publicado":
        return "0"
    if "(R$)" in col or "(R$/US$)" in col:
        return "#,##0.00"
    if "ao dia" in col:
        return "0.0000"
    if "(%" in col:
        return "0.00"
    return "#,##0"  # R$ milhões, pontos do Ibovespa, milhões de ações


with pd.ExcelWriter(PLANILHA, engine="openpyxl") as w:
    for nome, df in abas.items():
        df.to_excel(w, sheet_name=nome, index=False)
        ws = w.sheets[nome]
        ws.row_dimensions[1].height = 36
        for i, col in enumerate(df.columns, 1):
            cab = ws.cell(1, i)
            cab.fill, cab.font, cab.alignment, cab.border = CAB_FUNDO, CAB_FONTE, CAB_ALIN, CAB_BORDA
            ws.column_dimensions[cab.column_letter].width = 15
            fmt = formato(col, df[col])
            for (cel,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cel.number_format = fmt
                cel.alignment = CENTRO
for nome, df in abas.items():
    print(f"{nome:20s} {df.shape}")
print(f"\nGravado: {PLANILHA} ({os.path.getsize(PLANILHA) / 1e6:.2f} MB)")
